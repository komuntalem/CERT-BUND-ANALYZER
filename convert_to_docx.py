import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

md_path = r"C:\Users\Administrator\.gemini\antigravity\brain\54654864-0f78-4552-8faa-fb41f02bc721\CERT_BUND_GUI_Doc_Revised.md"
docx_path = r"C:\Users\Administrator\Desktop\INTERNSHIP\CERT-BUND-ANALYZER\CERT_BUND_GUI_Documentation_v2.docx"

if os.path.exists(docx_path):
    try:
        os.remove(docx_path)
    except Exception:
        pass

doc = Document()

with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_table = False
table_data = []

def process_table(doc, data):
    if not data:
        return
    rows = len(data)
    cols = max(len(row) for row in data)
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    for r_idx, row in enumerate(data):
        for c_idx, cell in enumerate(row):
            if c_idx < len(table.columns):
                # remove bold and code formatting
                cell_text = cell.strip()
                cell_text = cell_text.replace('**', '')
                cell_text = cell_text.replace('`', '')
                table.cell(r_idx, c_idx).text = cell_text
    doc.add_paragraph()

# Let's keep track if we are in a code block for the illustration
in_code_block = False

for line in lines:
    line = line.strip('\n')
    
    if line.startswith('```'):
        in_code_block = not in_code_block
        continue
        
    if in_code_block:
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Inches(0.5)
        for run in p.runs:
            run.font.name = 'Courier New'
        continue

    # Table logic
    if line.strip().startswith('|'):
        if not in_table:
            in_table = True
            table_data = []
        if re.match(r'^\|[-\s|]+\|$', line.strip()):
            continue
        row = [cell.strip() for cell in line.strip().strip('|').split('|')]
        table_data.append(row)
        continue
    else:
        if in_table:
            process_table(doc, table_data)
            in_table = False
            table_data = []
            
    if not line.strip():
        continue
        
    if line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=2)
    elif line.startswith('#### '):
        doc.add_heading(line[5:].strip(), level=3)
    elif line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=0)
    elif line.startswith('- '):
        clean = line[2:].replace('**', '').replace('`', '').replace('>', '')
        p = doc.add_paragraph(clean, style='List Bullet')
    elif re.match(r'^\d+\.\s+', line):
        clean = re.sub(r'^\d+\.\s+', '', line).replace('**', '').replace('`', '').replace('>', '')
        p = doc.add_paragraph(clean, style='List Number')
    elif line.startswith('---'):
        doc.add_paragraph('-' * 20)
    else:
        clean = line.replace('**', '').replace('`', '').replace('>', '')
        p = doc.add_paragraph(clean)

if in_table:
    process_table(doc, table_data)

doc.save(docx_path)
print("Saved to", docx_path)
